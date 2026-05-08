from pathlib import Path


class UnsupportedFileType(Exception):
    pass


def extract_text(filepath: Path) -> str:
    filepath = Path(filepath)
    suffix = filepath.suffix.lower()

    if suffix in ('.md', '.txt'):
        return filepath.read_text(encoding='utf-8')
    elif suffix == '.pdf':
        return _extract_pdf(filepath)
    elif suffix == '.docx':
        return _extract_docx(filepath)
    elif suffix == '.pptx':
        return _extract_pptx(filepath)
    elif suffix == '.xlsx':
        return _extract_xlsx(filepath)
    else:
        raise UnsupportedFileType(f"Unsupported file type: {suffix}")


def _extract_pdf(filepath: Path) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return '\n\n'.join(pages)


def _extract_docx(filepath: Path) -> str:
    from docx import Document
    doc = Document(filepath)
    return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(filepath: Path) -> str:
    from pptx import Presentation
    prs = Presentation(filepath)
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                lines.append(shape.text)
    return '\n'.join(lines)


def _extract_xlsx(filepath: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    lines = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f'Sheet: {sheet}')
        for row in ws.iter_rows(values_only=True):
            row_text = '\t'.join(str(c) if c is not None else '' for c in row)
            if row_text.strip():
                lines.append(row_text)
    return '\n'.join(lines)
